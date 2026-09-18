import io
import unittest

from openpyxl import load_workbook

from qa_testgen.infrastructure import work_item_batch as wib

TIPOS = [{"name": "User Story"}, {"name": "Task"}, {"name": "Bug"}]
CATALOGO = {
    "System.Description": {"type": "html", "read_only": False},
    "Microsoft.VSTS.Common.Priority": {"type": "integer", "read_only": False},
    "Microsoft.VSTS.Common.Severity": {"type": "string", "read_only": False},
    "System.Id": {"type": "integer", "read_only": True},
}
CAMPOS_POR_TIPO = {
    "User Story": [{"reference_name": "System.Title", "name": "Title", "always_required": True, "allowed_values": []}],
    "Task": [{"reference_name": "Microsoft.VSTS.Common.Priority", "name": "Priority", "always_required": True, "allowed_values": ["1", "2", "3", "4"]}],
    "Bug": [
        {"reference_name": "Microsoft.VSTS.Common.Severity", "name": "Severity", "always_required": True,
         "allowed_values": ["1 - Critical", "2 - High", "3 - Medium", "4 - Low"]},
        {"reference_name": "System.Id", "name": "ID", "always_required": True, "allowed_values": []},
    ],
}
WIDGET = {"System.Title", "System.Description", "System.AreaPath", "System.IterationPath", "System.AssignedTo", "System.Tags", "System.State"}
AREAS = ["Proj", "Proj\\Backend"]
ITERS = ["Proj\\Sprint 1"]
PESSOAS = [{"display_name": "Mikael Braga", "unique_name": "mikael.braga@refuturiza.com.br"}]


def _extras():
    return wib.campos_extras_por_tipo(TIPOS, CAMPOS_POR_TIPO, CATALOGO, WIDGET)


class TemplateTests(unittest.TestCase):
    def test_extras_ignore_widget_and_readonly_fields(self):
        extras = _extras()
        self.assertEqual([c["name"] for c in extras["User Story"]], [])
        self.assertEqual([c["name"] for c in extras["Task"]], ["Priority"])
        self.assertEqual([c["name"] for c in extras["Bug"]], ["Severity"])   # System.Id é somente-leitura

    def test_xlsx_template_has_sheets_headers_and_lists(self):
        data = wib.gerar_modelo_xlsx("Proj", TIPOS, _extras(), AREAS, ITERS, PESSOAS)
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Work Items", "Instruções", "Listas"])
        cab = [c.value for c in wb["Work Items"][1]]
        self.assertEqual(cab[:3], ["Ref", "Tipo *", "Título *"])
        self.assertIn("Priority", cab)
        self.assertIn("Severity", cab)
        exemplo = [c.value for c in wb["Work Items"][3]]
        self.assertEqual(exemplo[cab.index("Pai")], "#US1")   # Task filha da US1 do exemplo
        listas = [c.value for c in wb["Listas"][2]]
        self.assertEqual(listas[0], "User Story")
        self.assertEqual(listas[3], PESSOAS[0]["unique_name"])

    def test_csv_template_roundtrips_through_reader(self):
        csv_txt = wib.gerar_modelo_csv(TIPOS, _extras(), "Proj")
        linhas = wib.ler_arquivo("modelo.csv", csv_txt.encode("utf-8"))
        self.assertEqual(len(linhas), 1)
        self.assertEqual(linhas[0]["Tipo *"], "User Story")


class ReadAndValidateTests(unittest.TestCase):
    def _validar(self, linhas):
        cols = wib.colunas_do_modelo(TIPOS, _extras())
        return wib.validar_linhas(linhas, cols, TIPOS, _extras(), CATALOGO, AREAS, ITERS, PESSOAS)

    def test_reads_tab_txt_and_xlsx(self):
        txt = "Tipo\tTítulo\tPai\nTask\tFazer X\t7040\n".encode("utf-8")
        self.assertEqual(wib.ler_arquivo("a.txt", txt)[0]["Título"], "Fazer X")
        data = wib.gerar_modelo_xlsx("Proj", TIPOS, _extras(), AREAS, ITERS, PESSOAS)
        linhas = wib.ler_arquivo("modelo.xlsx", data)
        self.assertEqual(len(linhas), 2)
        self.assertEqual(linhas[1]["Pai"], "#US1")

    def test_validation_maps_fields_and_reports_errors(self):
        linhas = [
            {"Ref": "US1", "Tipo *": "user story", "Título *": "Login", "Descrição": "a\nb", "Area Path": "proj\\backend",
             "Iteration": "Proj\\Sprint 1", "Tags": "login; api", "Atribuído a (e-mail)": "Mikael Braga", "Pai": "", "Estado": "New", "Priority": ""},
            {"Ref": "T1", "Tipo *": "Task", "Título *": "Endpoint", "Pai": "#US1", "Priority": "2"},
            {"Ref": "", "Tipo *": "Epico", "Título *": "", "Pai": "abc", "Area Path": "Nada", "Atribuído a (e-mail)": "x@y.z"},
            {"Ref": "T2", "Tipo *": "Task", "Título *": "Sem prioridade", "Pai": "#NAOEXISTE", "Priority": "9"},
        ]
        r = self._validar(linhas)
        self.assertEqual(r[0]["erros"], [])
        self.assertEqual(r[0]["tipo"], "User Story")
        self.assertEqual(r[0]["campos"]["System.AreaPath"], "Proj\\Backend")
        self.assertEqual(r[0]["campos"]["System.AssignedTo"], PESSOAS[0]["unique_name"])
        self.assertEqual(r[0]["campos"]["System.Description"], "a<br>b")
        self.assertEqual(r[0]["campos"]["System.Tags"], "login; api")
        self.assertEqual(r[0]["state"], "New")
        self.assertEqual(r[1]["parent_ref"], "US1")
        self.assertEqual(r[1]["campos"]["Microsoft.VSTS.Common.Priority"], 2)
        self.assertEqual(r[1]["erros"], [])
        erros3 = " ".join(r[2]["erros"])
        for trecho in ("Tipo 'Epico'", "Título vazio", "Pai 'abc'", "Area Path 'Nada'", "Pessoa 'x@y.z'"):
            self.assertIn(trecho, erros3)
        erros4 = " ".join(r[3]["erros"])
        self.assertIn("valores aceitos", erros4)
        self.assertIn("#NAOEXISTE", erros4)

    def test_parent_ref_ordering(self):
        itens = [
            {"ref": "T1", "parent_ref": "US1"}, {"ref": "US1", "parent_ref": None},
            {"ref": "T2", "parent_ref": "T1"}, {"ref": "", "parent_ref": None},
        ]
        ordem = [i["ref"] for i in wib.ordenar_para_criacao(itens)]
        self.assertLess(ordem.index("US1"), ordem.index("T1"))
        self.assertLess(ordem.index("T1"), ordem.index("T2"))
        self.assertEqual(len(ordem), 4)


if __name__ == "__main__":
    unittest.main()
